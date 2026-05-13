import io
from typing import Optional
import pdfplumber
from pypdf import PdfReader
import openpyxl
import docx


# -----------------------------
#  PDF
# -----------------------------
def extract_from_pdf(content: bytes) -> str:
    """
    Extrae texto de un archivo PDF.
    Se usa pdfplumber porque maneja mejor columnas, tablas y layouts complejos.
    """
    text = ""

    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception:
        # Fallback con pypdf si pdfplumber falla
        reader = PdfReader(io.BytesIO(content))
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"

    return text.strip()


# -----------------------------
#  EXCEL
# -----------------------------
def extract_from_excel(content: bytes) -> str:
    """
    Extrae texto de un archivo Excel (.xlsx).
    Convierte cada fila en una línea de texto.
    """
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    text_parts = []

    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_values = [str(cell) for cell in row if cell is not None]
            if row_values:
                text_parts.append(" | ".join(row_values))

    return "\n".join(text_parts)


# -----------------------------
#  WORD
# -----------------------------
def extract_from_word(content: bytes) -> str:
    """
    Extrae texto de un archivo Word (.docx).
    """
    doc = docx.Document(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


# -----------------------------
#  TXT
# -----------------------------
def extract_from_text(content: bytes) -> str:
    """
    Extrae texto de un archivo de texto plano.
    Intenta UTF-8 y si falla usa Latin-1.
    """
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1", errors="ignore")


# -----------------------------
#  FUNCIÓN PRINCIPAL
# -----------------------------
def extract_text(content: bytes, content_type: Optional[str], filename: str) -> str:
    """
    Detecta el tipo de archivo y llama al extractor adecuado.
    Devuelve siempre texto plano.
    """
    filename_lower = filename.lower()

    # PDF
    if filename_lower.endswith(".pdf") or (content_type and "pdf" in content_type):
        return extract_from_pdf(content)

    # Excel
    if filename_lower.endswith(".xlsx") or filename_lower.endswith(".xls") or (content_type and "excel" in content_type):
        return extract_from_excel(content)

    # Word
    if filename_lower.endswith(".docx") or (content_type and "word" in content_type):
        return extract_from_word(content)

    # TXT
    if filename_lower.endswith(".txt") or (content_type and "text" in content_type):
        return extract_from_text(content)

    # Fallback: intentar como texto plano
    return extract_from_text(content)
